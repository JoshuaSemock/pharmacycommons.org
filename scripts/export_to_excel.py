"""Export the Pharmacy Commons Supabase database into 3 Excel workbooks via the REST API.

Needs SUPABASE_SECRET_KEY in a .env file next to this script (the service_role / secret key).
Usage:  python export_to_excel.py            (all 3 workbooks)
        python export_to_excel.py 1 2        (only the listed subsets)
"""
import csv, datetime, decimal, json, os, sys, time, uuid
from pathlib import Path
from urllib.parse import quote

import requests
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

HERE = Path(__file__).resolve().parent
STAMP = "2026-09-23"
API = "https://nenwovhyrdcdkhxzjiiv.supabase.co/rest/v1"
PAGE = 1000
MAX_CELL = 32767
MAX_SHEET_ROWS = 1_000_000  # data rows per sheet (Excel limit is 1,048,576 incl. header)

# primary keys (from pg_index); None = materialized view without a key
PK = {
    "biological_targets": ["pcid"], "chemont_links": ["pcid", "chemont_id"],
    "chemont_terms": ["chemont_id"], "class_agent_review": ["member_pcid", "class_slug"],
    "class_members": ["class_pcid", "member_pcid"],
    "class_rules": ["code"], "classyfire_lookup": ["pcid"],
    "clinical_concepts": ["pcid"], "clinical_statements": ["statement_id"],
    "combinations": ["pcid"], "drug_classes": ["pcid"], "entities": ["pcid"],
    "entity_brand_names": None, "entity_guidelines": ["pcid", "guideline_id"],
    "entity_label_rank": None, "fda_applications": ["appl_no"], "fda_products": ["product_id"],
    "formulations": ["pcid"], "functional_groups": ["pcid"], "guidelines": ["guideline_id"],
    "ingest_tokens": ["name"], "ingredient_moiety_map": None,
    "label_document_formulations": ["setid", "pcid"], "label_documents": ["setid"],
    "label_fetch_log": ["setid"], "label_section_text": ["setid", "section_key"],
    "label_sections": ["id"], "measurements": ["pcid"], "moieties": ["pcid"],
    "moiety_hierarchy": None, "pcid_blocks": ["block_id"], "pcid_retired": ["pcid"],
    "physiochemical": ["pcid"], "precise_forms": ["pcid"], "predicates": ["predicate"],
    "provider_verifications": ["user_id"], "revisions": ["id"],
    "rxclass_classes": ["class_id", "rela_source"],
    "rxclass_drug": ["pcid", "class_id", "rela_source", "rela"], "rxclass_lookup": ["pcid"],
    "rxclass_va_products": ["class_id", "product"], "rxnorm_brands": ["pcid", "brand_rxcui"],
    "rxnorm_lookup": ["pcid"], "saved_entities": ["user_id", "pcid_code"],
}

SUBSETS = {
    1: ("1_Core_Entities", [
        "entities", "moieties", "precise_forms", "formulations", "combinations",
        "pcid_blocks", "pcid_retired", "moiety_hierarchy", "ingredient_moiety_map",
        "entity_brand_names", "clinical_statements", "clinical_concepts", "predicates",
        "measurements", "biological_targets", "physiochemical", "functional_groups",
        "guidelines", "entity_guidelines", "revisions", "saved_entities",
        "provider_verifications", "ingest_tokens"]),
    2: ("2_Classification_RxNorm", [
        "drug_classes", "class_members", "chemont_terms", "chemont_links",
        "classyfire_lookup", "rxclass_classes", "rxclass_drug", "rxclass_lookup",
        "rxclass_va_products", "rxnorm_brands", "rxnorm_lookup", "class_rules",
        "class_agent_review"]),
    3: ("3_FDA_Labels", [
        "fda_applications", "fda_products", "label_documents",
        "label_document_formulations", "label_section_text", "label_fetch_log",
        "entity_label_rank", "label_sections"]),
}


def log(msg):
    line = f"[{datetime.datetime.now():%H:%M:%S}] {msg}"
    print(line, flush=True)
    with open(HERE / "export_log.txt", "a", encoding="utf-8") as f:
        f.write(line + "\n")


def load_key():
    env = HERE / ".env"
    if env.exists():
        for raw in env.read_text(encoding="utf-8").splitlines():
            if raw.strip().startswith("SUPABASE_SECRET_KEY="):
                val = raw.split("=", 1)[1].strip().strip('"').strip("'")
                if val and "PASTE" not in val:
                    return val
    key = os.environ.get("SUPABASE_SECRET_KEY")
    if not key:
        sys.exit("SUPABASE_SECRET_KEY not set: paste your secret key into .env")
    return key


def make_session(key):
    s = requests.Session()
    s.headers["apikey"] = key
    if key.startswith("eyJ"):  # legacy service_role JWT
        s.headers["Authorization"] = f"Bearer {key}"
    s.headers["Accept"] = "application/json"
    return s


def get(s, url, **kw):
    for attempt in range(6):
        try:
            r = s.get(url, timeout=120, **kw)
            if r.status_code in (200, 206):
                return r
            if r.status_code in (401, 403):
                sys.exit(f"Key rejected ({r.status_code}): {r.text[:300]}")
            err = f"HTTP {r.status_code}: {r.text[:300]}"
        except requests.RequestException as e:
            err = str(e)
        wait = 2 ** attempt
        log(f"    retry {attempt + 1} in {wait}s ({err})")
        time.sleep(wait)
    sys.exit(f"Giving up on {url}")


def row_count(s, table):
    r = get(s, f"{API}/{table}?select=*&limit=1", headers={"Prefer": "count=exact"})
    cr = r.headers.get("Content-Range", "*/0")
    return int(cr.split("/")[-1]) if cr.split("/")[-1].isdigit() else -1


def to_cell(v):
    """Convert a JSON value to something openpyxl accepts. Returns (value, full_text_if_truncated)."""
    if v is None or isinstance(v, (bool, int, float)):
        return v, None
    if isinstance(v, (dict, list)):
        s = json.dumps(v, ensure_ascii=False)
    else:
        s = str(v)
    s = ILLEGAL_CHARACTERS_RE.sub("", s)
    if len(s) > MAX_CELL:
        return s[:MAX_CELL], s
    return s, None


def iter_rows(s, table):
    """Yield every row of a table, page by page, in primary-key order."""
    pk = PK.get(table)
    if pk and len(pk) == 1:  # keyset pagination: fast even for millions of rows
        col, last = pk[0], None
        while True:
            url = f"{API}/{table}?select=*&order={col}.asc&limit={PAGE}"
            if last is not None:
                url += f"&{col}=gt.{quote(str(last), safe='')}"
            rows = get(s, url).json()
            if not rows:
                return
            yield from rows
            last = rows[-1][col]
    else:  # composite key or no key: offset pagination with a stable order
        if pk:
            order = ",".join(f"{c}.asc" for c in pk)
        else:
            probe = get(s, f"{API}/{table}?select=*&limit=1").json()
            if not probe:
                return
            order = ",".join(f"{c}.asc" for c in probe[0].keys())
        offset = 0
        while True:
            rows = get(s, f"{API}/{table}?select=*&order={order}&limit={PAGE}&offset={offset}").json()
            if not rows:
                return
            yield from rows
            offset += len(rows)


def safe_sheet_name(name, used):
    base = name[:31]
    cand, n = base, 2
    while cand.lower() in used:
        suffix = f"_{n}"
        cand = base[:31 - len(suffix)] + suffix
        n += 1
    used.add(cand.lower())
    return cand


def export_table(s, wb, table, used_names, trunc_writer, index_rows):
    t0 = time.time()
    expected = row_count(s, table)
    state = {"ws": None, "part": 0, "in_sheet": 0, "written": 0, "truncated": 0, "headers": None}
    sheets = []

    def new_sheet():
        state["part"] += 1
        nm = table if state["part"] == 1 else f"{table}_{state['part']}"
        ws = wb.create_sheet(safe_sheet_name(nm, used_names))
        ws.append(state["headers"])
        sheets.append(ws.title)
        state["ws"], state["in_sheet"] = ws, 0

    for row in iter_rows(s, table):
        if state["headers"] is None:
            state["headers"] = list(row.keys())
            new_sheet()
        elif state["in_sheet"] >= MAX_SHEET_ROWS:
            new_sheet()
        out = []
        for col in state["headers"]:
            val, full = to_cell(row.get(col))
            if full is not None:
                state["truncated"] += 1
                trunc_writer.writerow([table, state["ws"].title, state["in_sheet"] + 2, col, len(full), full])
            out.append(val)
        state["ws"].append(out)
        state["in_sheet"] += 1
        state["written"] += 1
        if state["written"] % 50_000 == 0:
            log(f"    {table}: {state['written']:,}/{expected:,} rows")

    if state["headers"] is None:  # empty table
        ws = wb.create_sheet(safe_sheet_name(table, used_names))
        ws.append(["(empty table)"])
        sheets.append(ws.title)

    w = state["written"]
    note = "" if w == expected else f"row count mismatch (database reported {expected:,})"
    ncols = len(state["headers"] or [])
    index_rows.append([table, ", ".join(sheets), ncols, w, state["truncated"], note])
    log(f"  {table}: {w:,} rows, {state['truncated']} truncated cells, {time.time() - t0:.0f}s"
        + (f"  !! {note}" if note else ""))


def export_subset(s, num):
    name, tables = SUBSETS[num]
    out_path = HERE / f"PharmacyCommons_{name}_{STAMP}.xlsx"
    trunc_path = HERE / f"PharmacyCommons_{name}_{STAMP}_truncated_cells.csv"
    log(f"Workbook {num}: {out_path.name}")
    wb = Workbook(write_only=True)
    used, index_rows = {"index"}, []
    with open(trunc_path, "w", newline="", encoding="utf-8") as tf:
        tw = csv.writer(tf)
        tw.writerow(["table", "sheet", "excel_row", "column", "full_length", "full_value"])
        for t in tables:
            export_table(s, wb, t, used, tw, index_rows)
    idx = wb.create_sheet("Index")
    idx.append([f"Pharmacy Commons Supabase backup ({name}), exported {datetime.datetime.now():%Y-%m-%d %H:%M}"])
    idx.append(["Cells longer than 32,767 characters are cut off in the sheets; their full text is in "
                + trunc_path.name])
    idx.append([])
    idx.append(["table", "sheet(s)", "columns", "rows", "truncated cells", "note"])
    for r in index_rows:
        idx.append(r)
    wb._sheets.insert(0, wb._sheets.pop())  # move Index to the front
    log(f"  saving {out_path.name} ...")
    wb.save(out_path)
    if not any(r[4] for r in index_rows):
        trunc_path.unlink()
    log(f"  saved ({out_path.stat().st_size / 1e6:,.1f} MB)")


def main():
    wanted = [int(a) for a in sys.argv[1:]] or [1, 2, 3]
    s = make_session(load_key())
    get(s, f"{API}/pcid_blocks?select=*&limit=1")  # fails fast on a bad key
    log(f"Key accepted. Exporting subsets {wanted}")
    for n in wanted:
        export_subset(s, n)
    log("Done.")


if __name__ == "__main__":
    main()
